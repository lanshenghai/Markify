from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def xlsx_to_markdown(path: str | Path, max_rows: int | None = 5000) -> str:
    """Convert each worksheet to a Markdown section with a pipe table."""
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in wb.worksheets:
            title = f"## {sheet.title}"
            rows_iter = sheet.iter_rows(values_only=True)
            rows: list[tuple] = []
            for i, row in enumerate(rows_iter):
                if max_rows is not None and i >= max_rows:
                    rows.append(tuple([f"... (truncated after {max_rows} rows)"]))
                    break
                rows.append(tuple("" if c is None else str(c) for c in row))
            if not rows:
                parts.append(f"{title}\n\n_(empty)_")
                continue
            max_cols = max(len(r) for r in rows)
            norm = [tuple(list(r) + [""] * (max_cols - len(r))) for r in rows]
            header = norm[0]
            lines = [title, "", "| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
            for r in norm[1:]:
                lines.append("| " + " | ".join(str(c).replace("\n", " ") for c in r) + " |")
            parts.append("\n".join(lines))
        return "\n\n".join(parts)
    finally:
        wb.close()
