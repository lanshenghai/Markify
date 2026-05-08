from __future__ import annotations

import hashlib
import io
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_PROG_ID_EXT: dict[str, str] = {
    "Excel.Sheet.12": ".xlsx",
    "Excel.Sheet.8": ".xls",
    "Word.Document.12": ".docx",
    "Word.Document.8": ".doc",
    "PowerPoint.Show.12": ".pptx",
    "PowerPoint.Show.8": ".ppt",
    "Package": ".bin",
    "AcroExch.Document.DC": ".pdf",
}


def _safe_filename(name: str, fallback: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name.strip())
    return name or fallback


def _extension_for_prog_id(prog_id: str | None) -> str:
    if not prog_id:
        return ".bin"
    return _PROG_ID_EXT.get(prog_id, ".bin")


def _collect_text_from_text_frame(shape: Any) -> list[str]:
    lines: list[str] = []
    if not shape.has_text_frame:
        return lines
    tf = shape.text_frame
    for para in tf.paragraphs:
        parts = [run.text for run in para.runs]
        text = "".join(parts).strip()
        if text:
            lines.append(text)
    return lines


def _table_to_markdown(shape: Any) -> str | None:
    if not shape.has_table:
        return None
    rows: list[list[str]] = []
    for row in shape.table.rows:
        rows.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
    if not rows:
        return None
    w = len(rows[0])
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in range(w)) + " |",
    ]
    for r in rows[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _walk_shapes(
    shapes: Any,
    text_blocks: list[str],
    ole_writer: Callable[[Any], None] | None,
) -> None:
    for shape in shapes:
        st = shape.shape_type
        if st == MSO_SHAPE_TYPE.GROUP:
            _walk_shapes(shape.shapes, text_blocks, ole_writer)
            continue

        if shape.has_table:
            md = _table_to_markdown(shape)
            if md:
                text_blocks.append(md)
            continue

        if shape.has_text_frame:
            for line in _collect_text_from_text_frame(shape):
                text_blocks.append(line)

        if ole_writer is not None and st in (
            MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
            MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        ):
            ole_writer(shape)


def _walk_picture_shapes(shapes: Any, blobs: list[bytes]) -> None:
    """Collect raster image bytes from Picture shapes (often duplicates of ppt/media)."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk_picture_shapes(shape.shapes, blobs)
            continue
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            image = getattr(shape, "image", None)
            if image is not None and getattr(image, "blob", None):
                blobs.append(image.blob)
        except (AttributeError, ValueError):
            continue


def _ocr_image_bytes(data: bytes, lang: str) -> str:
    from PIL import Image
    import pytesseract

    img = Image.open(io.BytesIO(data))
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    return pytesseract.image_to_string(img, lang=lang).strip()


def _extract_zip_embeddings(
    pptx_path: Path,
    out_dir: Path,
    written: dict[str, Path],
    counter: list[int],
) -> list[tuple[str, Path]]:
    """Copy everything under ppt/embeddings/ from the package."""
    saved: list[tuple[str, Path]] = []
    with zipfile.ZipFile(pptx_path, "r") as zf:
        for name in zf.namelist():
            if not name.startswith("ppt/embeddings/") or name.endswith("/"):
                continue
            rel = Path(name).name
            data = zf.read(name)
            key = name
            if key in written:
                continue
            dest = out_dir / rel
            dest.parent.mkdir(parents=True, exist_ok=True)
            if dest.exists():
                counter[0] += 1
                stem = dest.stem
                dest = dest.with_name(f"{stem}_{counter[0]}{dest.suffix}")
            dest.write_bytes(data)
            written[key] = dest
            saved.append((name, dest))
    return saved


def _save_ole_blob(
    shape: Any,
    out_dir: Path,
    written_names: set[str],
    counter: list[int],
) -> Path | None:
    ole = getattr(shape, "ole_format", None)
    if ole is None:
        return None
    try:
        blob = ole.blob
        prog_id = ole.prog_id
    except (AttributeError, ValueError, NotImplementedError):
        return None
    if not blob:
        return None
    ext = _extension_for_prog_id(prog_id)
    base = _safe_filename(getattr(shape, "name", "") or "embedded", "embedded")
    fname = f"{base}{ext}"
    if fname in written_names:
        counter[0] += 1
        fname = f"{Path(fname).stem}_{counter[0]}{ext}"
    written_names.add(fname)
    path = out_dir / fname
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(blob)
    return path


def _ocr_all_images(
    pptx_path: Path,
    prs: Presentation,
    attachments_dir: Path,
    lang: str,
    errors: list[str],
) -> list[tuple[str, str]]:
    """OCR images from ``ppt/media`` and Picture shapes; dedupe by SHA-256."""
    results: list[tuple[str, str]] = []
    seen_hashes: set[str] = set()
    media_dir = attachments_dir / "_media_copy"
    media_dir.mkdir(parents=True, exist_ok=True)
    image_ext = {".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".bmp", ".webp"}

    def consider(label: str, data: bytes) -> None:
        digest = hashlib.sha256(data).hexdigest()
        if digest in seen_hashes:
            return
        seen_hashes.add(digest)
        try:
            text = _ocr_image_bytes(data, lang)
        except Exception as e:  # noqa: BLE001
            errors.append(f"{label}: {e}")
            return
        if text:
            results.append((label, text))

    with zipfile.ZipFile(pptx_path, "r") as zf:
        for name in zf.namelist():
            if not name.startswith("ppt/media/") or name.endswith("/"):
                continue
            p = Path(name)
            if p.suffix.lower() not in image_ext:
                continue
            data = zf.read(name)
            dest = media_dir / p.name
            if dest.exists():
                dest = dest.with_stem(dest.stem + "_dup")
            dest.write_bytes(data)
            rel = str(dest.relative_to(attachments_dir)).replace("\\", "/")
            consider(rel, data)

    for si, slide in enumerate(prs.slides, start=1):
        blobs: list[bytes] = []
        _walk_picture_shapes(slide.shapes, blobs)
        for j, blob in enumerate(blobs, start=1):
            consider(f"picture_shape/slide{si}_#{j}", blob)

    return results


def pptx_to_markdown(
    path: str | Path,
    *,
    attachments_dir: Path | None = None,
    ocr_images: bool = True,
    ocr_lang: str = "chi_sim+eng",
) -> str:
    """
    Convert .pptx to Markdown.

    - Extracts text from grouped shapes, text boxes, and tables.
    - When ``attachments_dir`` is set, writes embedded package parts under
      ``ppt/embeddings/`` and OLE object blobs from shapes into that folder.
    - When ``ocr_images`` is True, runs Tesseract on images in ``ppt/media/``
      and on Picture shapes (requires Tesseract OCR installed on the system).
    """
    path = Path(path).resolve()
    prs = Presentation(path)

    ole_paths: list[Path] = []
    written_ole_names: set[str] = set()
    counter: list[int] = [0]

    def ole_writer(shape: Any) -> None:
        if attachments_dir is None:
            return
        emb_dir = attachments_dir / "embeddings"
        p = _save_ole_blob(shape, emb_dir, written_ole_names, counter)
        if p is not None:
            ole_paths.append(p)

    sections: list[str] = []
    errors: list[str] = []

    for i, slide in enumerate(prs.slides, start=1):
        blocks: list[str] = []
        _walk_shapes(slide.shapes, blocks, ole_writer)
        body = "\n\n".join(blocks) if blocks else ""
        sections.append(f"## Slide {i}\n\n{body}".strip())

    out = "\n\n".join(sections)

    zip_saved: list[tuple[str, Path]] = []
    if attachments_dir is not None:
        attachments_dir = Path(attachments_dir).resolve()
        attachments_dir.mkdir(parents=True, exist_ok=True)
        written: dict[str, Path] = {}
        zip_saved = _extract_zip_embeddings(path, attachments_dir / "embeddings", written, counter)

        attach_lines = ["### Attachments", ""]
        for arcname, dest in zip_saved:
            rel = dest.relative_to(attachments_dir)
            attach_lines.append(f"- `{rel.as_posix()}` ← `{arcname}`")
        for p in ole_paths:
            rel = p.relative_to(attachments_dir)
            attach_lines.append(f"- `{rel.as_posix()}` (OLE object)")
        if len(attach_lines) > 2:
            out = out + "\n\n" + "\n".join(attach_lines)

    if ocr_images:
        if attachments_dir is not None:
            ocr_root = Path(attachments_dir).resolve()
            ocr_root.mkdir(parents=True, exist_ok=True)
            ocr_results = _ocr_all_images(path, prs, ocr_root, ocr_lang, errors)
        else:
            with tempfile.TemporaryDirectory(prefix="markify_pptx_ocr_") as td:
                ocr_results = _ocr_all_images(path, prs, Path(td), ocr_lang, errors)
        if ocr_results:
            ocr_lines = ["### Images (OCR)", ""]
            for rel, text in ocr_results:
                ocr_lines.append(f"**{rel}**\n\n```\n{text}\n```")
            out = out + "\n\n" + "\n".join(ocr_lines)

    if errors:
        out += "\n\n### OCR warnings\n\n" + "\n".join(f"- {e}" for e in errors)

    return out
